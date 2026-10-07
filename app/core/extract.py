# app/core/extract.py

"""
Free text (skills typed by a student, a resume, a job description, a syllabus) -> ESCO skills.

Pipeline
  1. read      PDF (pdfplumber) / DOCX (python-docx) / plain text, in memory only
  2. language  Devanagari -> "hi", Gurmukhi -> "pa", common romanised-Hindi words -> "hi-Latn",
               otherwise "en"
  3. exact     longest-match n-gram lookup (1-6 words) over ESCO preferred / alternative /
               short labels and Naukri tags already linked to ESCO. Runs on Latin-script words
               in any language, because Hindi and Punjabi speakers mix in English skill names
  4. loanwords Hindi / Punjabi text: English terms written in Devanagari or Gurmukhi
               ("एक्सेल", "ਟੈਲੀ") are mapped to English with a curated lexicon
               (data/reference/loanwords_hi_pa.csv) and resolved by the English matcher.
               Multilingual embeddings link such transliterations unreliably, at high cosine
  5. embedding the remaining short phrases (list items, comma-separated chunks), with filler
               words trimmed, are linked by embedding similarity with the pipeline threshold:
               native-script phrases with the multilingual model, Latin-script ones with the
               English model and its single-word guard

Nothing is stored. Every linked skill carries the text span it came from, its score and
the method, so the UI can show *why* a skill was picked up.
"""

from __future__ import annotations

import io
import re
import threading
import unicodedata
from dataclasses import dataclass
from functools import lru_cache

import numpy as np
import pandas as pd

from app.core.settings import get_settings

MAX_UPLOAD_BYTES = 5 * 1024 * 1024
MAX_PHRASES = 400
ROMAN_HINDI = {"hai", "hain", "karna", "karta", "karti", "aata", "aati", "mujhe", "mein", "aur", "kaam", "ka", "ki",
               "ke", "seekha", "jaanta", "jaanti", "sakta", "sakti", "chalana", "banana", "likhna", "bolna", "nahi"}
# Filler words trimmed from the edges of a phrase before linking ("मुझे एक्सेल आता है" -> "एक्सेल").
HINDI_STOP = {"मुझे", "मैं", "मैंने", "मेरा", "मेरी", "मेरे", "हम", "को", "का", "की", "के", "में", "से", "पर", "या", "भी",
              "है", "हैं", "हूँ", "हूं", "था", "थी", "थे", "आता", "आती", "आते", "जानता", "जानती", "सकता", "सकती", "कर",
              "करना", "करता", "करती", "किया", "थोड़ी", "थोड़ा", "बहुत", "अच्छा", "अच्छी", "अनुभव", "साल", "वर्ष",
              "काम", "बात", "ज्ञान", "जानकारी"}
PUNJABI_STOP = {"ਮੈਨੂੰ", "ਮੈਂ", "ਮੇਰਾ", "ਮੇਰੀ", "ਤੇ", "ਦਾ", "ਦੀ", "ਦੇ", "ਵਿੱਚ", "ਹੈ", "ਹਾਂ", "ਹਨ", "ਆਉਂਦਾ", "ਆਉਂਦੀ", "ਕਰ",
                "ਸਕਦਾ", "ਸਕਦੀ", "ਤਜਰਬਾ", "ਸਾਲ", "ਕੰਮ", "ਥੋੜ੍ਹੀ", "ਬਹੁਤ", "ਨੂੰ", "ਨਾਲ", "ਗੱਲ", "ਜਾਣਕਾਰੀ"}
STOPWORDS = HINDI_STOP | PUNJABI_STOP | ROMAN_HINDI | {"thoda", "thodi", "bahut", "accha", "anubhav", "saal", "bhi", "se",
                                                         "hoon", "hu", "main", "mera", "meri", "ko", "aata", "aati", "aate"}
# Whitespace lookarounds, not \b: Indic vowel signs are not "word" characters, so \b misfires inside words.
_SPLIT = re.compile(r"[\n\r;,•·●▪◦|/।॥]+|\s+-\s+|(?<!\S)(?:और|ਅਤੇ|and|aur)(?!\S)", re.I)
_WORD = re.compile(r"[A-Za-z0-9+#.]+|[ऀ-ॿ]+|[਀-੿]+")
_EDGE = " .:()[]{}\"'।॥"
LOANWORD_MAX_WORDS = 3


def read_upload(filename: str, content: bytes) -> str:
    if len(content) > MAX_UPLOAD_BYTES:
        raise ValueError("File is larger than 5 MB")
    name = filename.lower()
    if name.endswith(".pdf") or content[:4] == b"%PDF":
        import pdfplumber

        with pdfplumber.open(io.BytesIO(content)) as pdf:
            return "\n".join((p.extract_text() or "") for p in pdf.pages[:40])
    if name.endswith(".docx"):
        import docx

        d = docx.Document(io.BytesIO(content))
        parts = [p.text for p in d.paragraphs]
        for table in d.tables:
            for row in table.rows:
                parts.extend(c.text for c in row.cells)
        return "\n".join(parts)
    if name.endswith((".txt", ".md", ".csv")) or not name:
        return content.decode("utf-8", errors="replace")
    raise ValueError("Unsupported file type: use PDF, DOCX or plain text")


def detect_language(text: str) -> str:
    dev = len(re.findall(r"[ऀ-ॿ]", text))
    gur = len(re.findall(r"[਀-੿]", text))
    if gur > 5 and gur >= dev:
        return "pa"
    if dev > 5:
        return "hi"
    words = re.findall(r"[a-z]+", text.lower())
    if words and sum(w in ROMAN_HINDI for w in words) / len(words) > 0.15:
        return "hi-Latn"
    return "en"


def trim_stopwords(phrase: str, stop: set[str]) -> str:
    """Drop filler words from both ends of a phrase (never from the middle)."""
    def filler(w: str) -> bool:
        w = w.strip(_EDGE).lower()
        return w in stop or w.isdigit()

    words = phrase.split()
    while words and filler(words[0]):
        words.pop(0)
    while words and filler(words[-1]):
        words.pop()
    return " ".join(words).strip(_EDGE)


def candidate_phrases(text: str, stop: set[str] | None = None) -> list[str]:
    out, seen = [], set()
    for chunk in _SPLIT.split(text):
        chunk = re.sub(r"\s+", " ", chunk).strip(_EDGE)
        if stop:
            chunk = trim_stopwords(chunk, stop)
        n_words = len(chunk.split())
        if 1 <= n_words <= 6 and 2 <= len(chunk) <= 60 and chunk.lower() not in seen:
            seen.add(chunk.lower())
            out.append(chunk)
    return out[:MAX_PHRASES]


@lru_cache(maxsize=1)
def load_loanwords() -> dict[str, str]:
    """Devanagari / Gurmukhi spelling of an English workplace term -> the English term."""
    path = get_settings().reference_data_dir / "loanwords_hi_pa.csv"
    if not path.exists():
        return {}
    df = pd.read_csv(path, comment="#", dtype=str, keep_default_na=False)
    return {unicodedata.normalize("NFC", t.strip()): e.strip() for t, e in zip(df["term"], df["english"], strict=True)}


def split_loanwords(phrase: str, lexicon: dict[str, str], stop: set[str] | None = None) -> tuple[list[tuple[str, str]], str]:
    """
    Longest-match lookup of transliterated English terms inside a phrase.
    Returns [(source span, English term)] and what is left of the phrase (filler trimmed).
    """
    words = unicodedata.normalize("NFC", phrase).split()
    hits, rest, i = [], [], 0
    while i < len(words):
        for n in range(min(LOANWORD_MAX_WORDS, len(words) - i), 0, -1):
            span = " ".join(w.strip(_EDGE) for w in words[i:i + n])
            if span in lexicon:
                hits.append((span, lexicon[span]))
                i += n
                break
        else:
            rest.append(words[i])
            i += 1
    residual = " ".join(rest)
    return hits, trim_stopwords(residual, stop) if stop else residual


@dataclass
class LinkedSkill:
    uri: str
    label: str
    score: float
    method: str       # exact | transliteration | embedding
    source_text: str

    def as_dict(self) -> dict:
        return {"uri": self.uri, "label": self.label, "score": round(self.score, 4), "method": self.method,
                "source_text": self.source_text}


class SkillExtractor:
    """Lazy singleton: builds the label dictionary and embedding linkers on first use."""

    _instance: SkillExtractor | None = None
    _lock = threading.Lock()

    @classmethod
    def get(cls) -> SkillExtractor:
        with cls._lock:
            if cls._instance is None:
                cls._instance = cls()
            return cls._instance

    def __init__(self):
        from ml_pipeline.india.link import esco_concepts, label_variants
        from ml_pipeline.india.run import load_thresholds

        s = get_settings()
        self.concepts = esco_concepts(s.esco_data_dir, "skill")
        self.label_of = dict(zip(self.concepts["uri"], self.concepts["preferredLabel"], strict=True))
        th = load_thresholds()
        self.threshold = float(th["skill_link_threshold"])
        self.single_token_threshold = th.get("single_token_threshold")
        self.threshold_status = th["status"]
        # Exact dictionary: ESCO labels (preferred > short > alt) plus Naukri tags already linked to ESCO.
        self.exact: dict[str, str] = {}
        priority: dict[str, int] = {}
        rank = {"preferred": 0, "preferred_short": 1, "alt": 2}
        for row in self.concepts.itertuples(index=False):
            for text, kind in label_variants(row.preferredLabel, list(row.altLabels)):
                key = self._norm(text)
                if key and (key not in priority or rank[kind] < priority[key]):
                    self.exact[key], priority[key] = row.uri, rank[kind]
        tags_path = s.processed_data_dir / "india" / "tag_links.parquet"
        if tags_path.exists():
            tags = pd.read_parquet(tags_path, columns=["tag", "uri", "accepted", "freq"])
            for r in tags[tags["accepted"] & (tags["freq"] >= 5)].itertuples(index=False):
                self.exact.setdefault(self._norm(r.tag), r.uri)
        self.max_n = 6
        self._linkers: dict[str, object] = {}

    @staticmethod
    def _norm(t: str) -> str:
        return re.sub(r"\s+", " ", str(t).lower().strip(" .,;:-_()[]"))

    def _linker(self, multilingual: bool):
        key = "multi" if multilingual else "en"
        if key not in self._linkers:
            from ml_pipeline.india.link import ConceptLinker, EmbeddingModel

            s = get_settings()
            name = s.multilingual_model_name if multilingual else s.model_name
            enc = EmbeddingModel(name)
            self._linkers[key] = ConceptLinker(self.concepts, enc, s.artifacts_dir / "cache" / "emb",
                                               f"esco_skill_{name.rsplit('/', 1)[-1]}")
        return self._linkers[key]

    def _exact_spans(self, text: str) -> tuple[list[LinkedSkill], set[int]]:
        """Longest-match n-gram lookup; returns hits and the token positions they cover."""
        tokens = [m.group(0) for m in _WORD.finditer(text)]
        lower = [t.lower() for t in tokens]
        hits, used = [], set()
        i = 0
        while i < len(tokens):
            for n in range(min(self.max_n, len(tokens) - i), 0, -1):
                key = " ".join(lower[i:i + n])
                uri = self.exact.get(key)
                if uri and not (n == 1 and len(key) <= 2 and key not in {"c", "r"}):
                    hits.append(LinkedSkill(uri, self.label_of.get(uri, key), 1.0, "exact", " ".join(tokens[i:i + n])))
                    used.update(range(i, i + n))
                    i += n
                    break
            else:
                i += 1
        return hits, used

    def extract(self, text: str, language: str | None = None, use_embeddings: bool = True) -> dict:
        text = text or ""
        lang = language or detect_language(text)
        found: dict[str, LinkedSkill] = {}
        # Latin-script skill names are matched exactly whatever the language of the sentence around them.
        for h in self._exact_spans(text)[0]:
            found.setdefault(h.uri, h)
        unlinked: list[str] = []
        english: list[tuple[str, str]] = []      # (text to link with the English model, span shown to the user)
        native: list[str] = []                   # phrases for the multilingual model
        if lang == "en":
            english = [(p, p) for p in candidate_phrases(text)]
        else:
            lexicon = load_loanwords()
            for phrase in candidate_phrases(text, STOPWORDS):
                hits, rest = split_loanwords(phrase, lexicon, STOPWORDS)
                for span, term in hits:
                    uri = self.exact.get(self._norm(term))
                    if uri:
                        found.setdefault(uri, LinkedSkill(uri, self.label_of.get(uri, term), 1.0, "transliteration", span))
                    else:
                        english.append((term, f"{span} ({term})"))
                if len(rest) >= 2 and rest.isascii():
                    english.append((rest, rest))
                elif len(rest) >= 2:
                    native.append(rest)
        covered = [self._norm(f.source_text) for f in found.values()]

        def adds_nothing(t: str) -> bool:
            """Already matched exactly, inside a matched span, or a matched span plus at most one generic word."""
            n = self._norm(t)
            if n in self.exact or any(n in c for c in covered):
                return True
            return any(c in n and len(n.replace(c, " ").split()) <= 1 for c in covered if c)

        english = [(t, src) for t, src in english if not adds_nothing(t)]
        if use_embeddings and english:
            res = self._linker(multilingual=False).link([t for t, _ in english], threshold=self.threshold,
                                                        single_token_threshold=self.single_token_threshold)
            for r, (_, src) in zip(res.itertuples(index=False), english, strict=True):
                if r.accepted and r.uri not in found:
                    found[r.uri] = LinkedSkill(r.uri, r.label, float(r.score), r.method, src)
                elif not r.accepted:
                    unlinked.append(src)
        if use_embeddings and native:
            res = self._linker(multilingual=True).link(native, threshold=self.multilingual_threshold)
            for r in res.itertuples(index=False):
                if r.accepted and r.uri not in found:
                    found[r.uri] = LinkedSkill(r.uri, r.label, float(r.score), r.method, r.text)
                elif not r.accepted:
                    unlinked.append(r.text)
        return {
            "language": lang,
            "skills": sorted((f.as_dict() for f in found.values()), key=lambda d: -d["score"]),
            "unlinked_phrases": unlinked[:50],
            "threshold": self.threshold if lang == "en" else self.multilingual_threshold,
            "threshold_status": self.threshold_status,
        }

    @property
    def multilingual_threshold(self) -> float:
        # Multilingual cosine scores run lower than English mpnet scores; until the team's
        # multilingual gold set tunes it, use the English threshold minus 0.1 (provisional).
        return max(0.5, self.threshold - 0.1)


def skill_set_from(uris: list[str] | None, text: str | None, language: str | None = None) -> tuple[list[str], dict | None]:
    """Resolve a request's skills: explicit ESCO URIs plus anything extracted from free text."""
    out = list(dict.fromkeys(uris or []))
    extraction = None
    if text and text.strip():
        extraction = SkillExtractor.get().extract(text, language)
        out += [s["uri"] for s in extraction["skills"] if s["uri"] not in out]
    return out, extraction


def top_k(scores: np.ndarray, k: int) -> np.ndarray:
    k = min(k, len(scores))
    if k <= 0:
        return np.array([], dtype=np.int64)
    part = np.argpartition(-scores, k - 1)[:k]
    return part[np.argsort(-scores[part], kind="stable")]
