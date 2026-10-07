"""
Unit tests for the Phase 2-6 building blocks: ranking metrics and splits, the serving
ranker, extraction helpers, salary calibration, effort, workforce fields, the vectorised
upskilling gains, and the report renderer.
"""

from __future__ import annotations

import io
import time
from pathlib import Path

import numpy as np
import pandas as pd
import pytest
import scipy.sparse as sp

FIXTURES = Path(__file__).parent / "fixtures"
ESCO_MINI = FIXTURES / "esco_mini"


# --- graph: metrics and splits ---------------------------------------------------------------------

def test_ranking_metrics_on_hand_computed_cases():
    from ml_pipeline.graph.metrics import ndcg_at_k, recall_at_k, reciprocal_rank

    ranked = np.array([5, 3, 9, 1])
    rel = {3: 1, 1: 1}
    assert recall_at_k(ranked, rel, k=2) == 0.5
    assert reciprocal_rank(ranked, rel) == 0.5
    dcg = 1 / np.log2(3) + 1 / np.log2(5)
    idcg = 1 + 1 / np.log2(3)
    assert ndcg_at_k(ranked, rel, k=10) == pytest.approx(dcg / idcg)
    assert ndcg_at_k(np.array([3, 1]), rel) == pytest.approx(1.0)
    assert recall_at_k(ranked, {}) == 0.0


def test_top_k_and_full_rank_respect_exclusions():
    from ml_pipeline.graph.metrics import full_rank_of, top_k_from_scores

    s = np.array([0.1, 0.9, 0.5, 0.9, 0.3])
    assert top_k_from_scores(s, 3, exclude=np.array([1])).tolist() == [3, 2, 4]
    # ties are broken pessimistically: both 0.9 items get rank 2
    assert full_rank_of(s, np.array([1, 3])).tolist() == [2, 2]
    assert full_rank_of(s, np.array([2]), exclude=np.array([1, 3])).tolist() == [1]


def test_paired_difference_detects_a_real_gap():
    from ml_pipeline.graph.metrics import paired_difference

    rng = np.random.default_rng(0)
    b = rng.random(500)
    assert paired_difference(b + 0.1, b)["significant"]
    assert not paired_difference(b, b)["significant"]


def test_graph_split_is_deterministic_and_close_to_70_10_10_10():
    from ml_pipeline.graph.data import split_of

    groups = [f"g{i}" for i in range(20000)]
    splits = pd.Series([split_of(g) for g in groups]).value_counts(normalize=True)
    assert splits["train"] == pytest.approx(0.7, abs=0.02)
    for s in ("val", "test", "pool"):
        assert splits[s] == pytest.approx(0.1, abs=0.02)
    assert [split_of(g) for g in groups[:50]] == [split_of(g) for g in groups[:50]]
    assert split_of("g1", seed=1) in {"train", "val", "test", "pool"}


# --- serving ranker ------------------------------------------------------------------------------------

@pytest.fixture
def small_R():
    return sp.csr_matrix(np.array([
        [1, 1, 0, 0],
        [1, 0, 1, 0],
        [0, 0, 1, 1],
        [1, 1, 1, 0],
    ]))


@pytest.mark.parametrize("kind", ["b1", "b3a"])
def test_job_ranker_prefers_overlap_and_supports_row_subsets(small_R, kind):
    from app.core.serving import JobRanker

    r = JobRanker(small_R, kind)
    s = r.score([0, 1])
    assert s[2] == 0 and s.argmax() in (0, 3)
    assert np.allclose(r.score([0, 1], np.array([2, 0])), s[[2, 0]])
    assert (r.score([]) == 0).all()


def test_occlusion_reports_each_skills_contribution(small_R):
    from app.core.serving import JobRanker

    r = JobRanker(small_R, "b3a")
    drop = r.occlusion([0, 3], row=0)
    assert drop[0] > 0 and drop[3] == 0  # skill 3 is not in posting 0


def test_model_choice_serves_winner_or_best_servable(monkeypatch):
    import ml_pipeline.common as common
    from app.core.serving import Serving

    def rep(winner):
        return {"comparison": {"winner": winner}, "models": {
            "b1": {"mean": {"t3": {"ndcg@10": 0.2}}}, "b3a": {"mean": {"t3": {"ndcg@10": 0.3}}},
            "hgt": {"mean": {"t3": {"ndcg@10": 0.4}}}}}

    monkeypatch.setattr(common, "read_report", lambda name: rep("b1"))
    assert Serving().model_choice["served"] == "b1"
    monkeypatch.setattr(common, "read_report", lambda name: rep("hgt"))
    choice = Serving().model_choice
    assert choice["served"] == "b3a" and choice["benchmark_winner"] == "hgt" and choice["note"]
    monkeypatch.setattr(common, "read_report", lambda name: None)
    assert Serving().model_choice["served"] == "b1"


# --- extraction helpers ---------------------------------------------------------------------------------

@pytest.mark.parametrize("text,lang", [
    ("I know Python and SQL", "en"),
    ("मुझे एक्सेल और टैली आता है", "hi"),
    ("ਮੈਨੂੰ ਐਕਸਲ ਅਤੇ ਟੈਲੀ ਆਉਂਦੀ ਹੈ", "pa"),
    ("mujhe excel aur tally aata hai", "hi-Latn"),
])
def test_detect_language(text, lang):
    from app.core.extract import detect_language

    assert detect_language(text) == lang


def test_candidate_phrases_splits_lists_and_dedupes():
    from app.core.extract import candidate_phrases

    out = candidate_phrases("Python, SQL; python\n• Data analysis and machine learning | Excel")
    assert out == ["Python", "SQL", "Data analysis", "machine learning", "Excel"]
    assert candidate_phrases("x" * 200) == []


def test_read_upload_text_docx_and_limits():
    import docx

    from app.core.extract import MAX_UPLOAD_BYTES, read_upload

    assert read_upload("cv.txt", b"Python, SQL") == "Python, SQL"
    d = docx.Document()
    d.add_paragraph("Skills: Java")
    t = d.add_table(rows=1, cols=1)
    t.cell(0, 0).text = "Tally"
    buf = io.BytesIO()
    d.save(buf)
    text = read_upload("cv.docx", buf.getvalue())
    assert "Skills: Java" in text and "Tally" in text
    with pytest.raises(ValueError, match="5 MB"):
        read_upload("big.txt", b"x" * (MAX_UPLOAD_BYTES + 1))
    with pytest.raises(ValueError, match="Unsupported"):
        read_upload("cv.exe", b"MZ")


def test_top_k_helper():
    from app.core.extract import top_k

    assert top_k(np.array([0.1, 0.5, 0.3]), 2).tolist() == [1, 2]
    assert top_k(np.array([0.1]), 5).tolist() == [0]
    assert top_k(np.array([]), 3).tolist() == []


# --- linking: single-token ambiguity guard ---------------------------------------------------------------

def test_single_token_guard_only_affects_one_word_embedding_links():
    from ml_pipeline.india.link import ConceptLinker
    from tests.test_india_link_gold import TrigramEncoder

    concepts = pd.DataFrame({"uri": ["s:dev", "s:py"], "preferredLabel": ["developmental psychology", "Python"],
                             "altLabels": [[], []]})
    linker = ConceptLinker(concepts, TrigramEncoder())
    texts = ["development", "python", "developmental psychologyy"]
    plain = linker.link(texts, threshold=0.3)
    guarded = linker.link(texts, threshold=0.3, single_token_threshold=0.99)
    assert plain.loc[0, "accepted"] and not guarded.loc[0, "accepted"]   # generic single word: rejected
    assert guarded.loc[1, "accepted"] and guarded.loc[1, "method"] == "exact"  # exact matches are never guarded
    assert guarded.loc[2, "accepted"] == plain.loc[2, "accepted"]          # multi-word text: unaffected


# --- salary -----------------------------------------------------------------------------------------------

def test_cqr_margin_reaches_target_coverage():
    from ml_pipeline.salary.model import cqr_margin

    rng = np.random.default_rng(0)
    y = rng.normal(size=4000)
    q = np.column_stack([np.full(4000, -0.5), np.zeros(4000), np.full(4000, 0.5)])  # too narrow on purpose
    m = cqr_margin(q[:2000], y[:2000], 0.8)
    assert m > 0
    cov = np.mean((y[2000:] >= q[2000:, 0] - m) & (y[2000:] <= q[2000:, 2] + m))
    assert cov == pytest.approx(0.8, abs=0.03)


def test_pinball_and_support_key():
    from ml_pipeline.salary.model import SalaryModel, pinball

    y = np.array([1.0, 2.0, 3.0])
    assert pinball(y, np.full(3, 2.0), 0.5) == pytest.approx(1 / 3)
    assert pinball(y, y, 0.9) == 0
    assert SalaryModel.support_key("2512", 2.0) == "2512|2"
    assert SalaryModel.support_key(None, float("nan")) == "NA|0"


# --- effort and workforce fields ------------------------------------------------------------------------------

def test_effort_model_on_esco_fixture():
    from ml_pipeline.upskilling.effort import NEAR, REUSE, EffortModel

    skills = pd.read_csv(ESCO_MINI / "skills_en.csv", dtype=str, keep_default_na=False)
    ids = skills["conceptUri"].drop_duplicates().tolist()
    em = EffortModel(ESCO_MINI, ids)
    base = em.costs(np.array([], dtype=np.int64))
    assert set(np.round(base, 2)) <= set(REUSE.values()) | {1.0}
    linked = next((i for i, n in enumerate(em.neighbours) if n), None)
    if linked is not None:
        nb = next(iter(em.neighbours[linked]))
        c = em.costs(np.array([linked]))
        assert c[nb] == pytest.approx(base[nb] * NEAR)
        assert em.explain(nb, np.array([linked]))["close_to_known_skills"] == [linked]


def test_skill_fields_and_posting_fields_on_fixture():
    from ml_pipeline.supply.workforce import UNASSIGNED, posting_fields, skill_fields

    fields = skill_fields(ESCO_MINI)
    assert all(isinstance(v, str) and v for v in fields.values())
    some = list(fields)[:2]
    jobs = pd.DataFrame({"skill_uris": [some, [], None]})
    out = posting_fields(jobs, fields)
    assert out.iloc[1] == UNASSIGNED and out.iloc[2] == UNASSIGNED
    if some:
        assert out.iloc[0] in set(fields.values())


def test_canonical_state_names():
    from ml_pipeline.supply.parse import canonical_state

    assert canonical_state("Jammu & Kashmir") == "Jammu and Kashmir"
    assert canonical_state("  Tamil   Nadu ") == "Tamil Nadu"


# --- upskilling: vectorised gains ------------------------------------------------------------------------------

@pytest.mark.parametrize("seed", range(5))
def test_vectorised_gains_match_scalar(seed):
    from tests.test_upskilling_engine import random_problem

    p = random_problem(seed, n_jobs=60, n_skills=16, salary=bool(seed % 2))
    cands = p.candidates()
    c = p.covered([int(cands[0])]) if len(cands) else p.base
    gf, gg = p.gains(c, cands)
    assert np.allclose(gf, [p._gain_F(int(s), c) for s in cands])
    assert np.allclose(gg, [p._gain_G(int(s), c) for s in cands])


def test_ilp_hint_is_never_worse_and_bound_is_valid():
    from ml_pipeline.upskilling.engine import brute_force, exact_ilp, frequency
    from tests.test_upskilling_engine import random_problem

    for seed in range(6):
        p = random_problem(seed)
        hint = frequency(p, 3).skills
        plan = exact_ilp(p, 3, time_limit=5, hint=hint)
        assert plan.F >= p.F(hint)
        best = brute_force(p, 3).F
        assert plan.bound is not None and plan.bound + 1e-6 >= best


def test_market_pool_filters():
    from tests.test_yojak_api import _market

    m = _market()
    assert len(m.pool()) == 90
    assert set(m.jobs.iloc[m.pool(tiers=[3])]["primary_tier"]) == {3.0}
    assert set(m.jobs.iloc[m.pool(isco2="24")]["occupation_uri"]) == {"o:acc"}
    assert len(m.pool(occupation_uri="o:dev", exp_bands=["0-1"])) == 30
    assert "g0" not in set(m.jobs.iloc[m.pool(exclude_group="g0")]["dup_group_id"])
    p = m.problem(m.pool(isco2="25"), np.array([0]), 0.6, value="salary")
    assert p.v.max() == 400000.0


# --- report rendering ------------------------------------------------------------------------------------------

def test_reporting_renders_pending_without_reports():
    from ml_pipeline.reporting import REPORTS, build_evaluation, readme_block, render_readme

    empty = dict.fromkeys(REPORTS)
    block = readme_block(empty)
    assert block.count("*Pending:*") == 5
    doc = build_evaluation(empty)
    assert "# Yojak: evaluation" in doc and "Limitations" in doc
    readme = "intro\n<!-- results:start -->\nold\n<!-- results:end -->\nfooter"
    out = render_readme(readme, empty)
    assert "old" not in out and out.startswith("intro") and out.endswith("footer")
    assert render_readme(out, empty) == out  # idempotent


def test_reporting_uses_real_numbers_when_reports_exist():
    from ml_pipeline.reporting import load_reports, readme_block

    reports = load_reports(Path(__file__).parents[1] / "reports")
    block = readme_block(reports)
    if reports["salary_eval"]:
        assert f"{reports['salary_eval']['test']['n']:,}" in block
    if reports["data_quality"]:
        assert f"{reports['data_quality']['dataset']['raw_rows']:,}" in block


def test_reporting_models_and_upskilling_sections():
    from ml_pipeline.reporting import impact_section, models_section, upskilling_section

    t = {"recall@10": 0.5, "ndcg@10": 0.4, "mrr": 0.3, "recall@10_ci95": [0.4, 0.6], "ndcg@10_ci95": [0.3, 0.5],
         "mrr_ci95": [0.2, 0.4]}
    mean = {"t1": t, "t3": t, "latency": {"jobs_for_profile": {"p50_ms": 3.0, "p95_ms": 5.0}},
            "memory": {"rss_mb": 900}}
    mc = {"models": {"b1": {"name": "B1 TF-IDF kNN", "mean": mean}, "hgt": {"name": "M HGT", "mean": mean}},
          "comparison": {"winner": "b1", "winner_name": "B1 TF-IDF kNN", "rule": "highest T3 NDCG@10",
                         "hgt_vs_best_baseline": {"baseline": "B1", "t3": {"mean_diff": -0.01, "ci95": [-0.02, 0.0]}},
                         "hgt_beats_best_baseline_on_t3": False}}
    out = models_section(mc)
    assert "B1 TF-IDF kNN" in out and "does not beat" in out and "n/a" in out  # no T2 -> n/a
    s = {"mean_F": 5.0, "mean_gap_pct": 1.0, "p95_gap_pct": 3.0, "share_optimal": 0.9, "runtime_ms_p50": 1.0,
         "runtime_ms_p95": 2.0}
    grid = {"k=3,tau=0.6": {"instances": 10, "mean_optimum": 5.2, "ilp_proven_optimal_share": 1.0,
                            **{m: s for m in ("exact_ilp", "greedy_F", "lazy_greedy_G", "frequency")}}}
    assert "Postings unlocked" in upskilling_section({"grid": grid})
    im = {"users": 3, "users_by_tier": {"2": 2, "3": 1}, "eligible_before_median": 1,
          "eligible_after_frequency_median": 2, "eligible_after_optimal_median": 4,
          "extra_postings_vs_frequency": {"mean": 2.0, "median": 2.0, "iqr": [1.0, 3.0]},
          "share_users_with_more_postings": 0.66, "share_users_no_worse": 1.0,
          "assumptions": {"k": 3, "eligibility": "e", "person": "p", "pool": "q", "no_causal_claim": "no claim"},
          "illustration": {"label": "ILLUSTRATION", "per_1000_users_extra_reachable_postings": 2000}}
    assert "+2.00" in impact_section(im)


# --- impact estimate ------------------------------------------------------------------------------------------

def test_impact_estimate_on_small_market(monkeypatch):
    import ml_pipeline.impact as impact
    from tests.test_yojak_api import _market

    monkeypatch.setattr(impact, "split_of", lambda g: "test")  # every posting may seed a pseudo-user
    body = impact.estimate(_market(), n_users=8, min_pool=5, min_skills=2, ilp_seconds=1)
    assert 0 < body["users"] <= 8
    assert body["users_by_tier"]["2"] + body["users_by_tier"]["3"] == body["users"]
    # the optimal plan is never worse than frequency advice
    assert body["share_users_no_worse"] == 1.0
    assert body["eligible_after_optimal_median"] >= body["eligible_after_frequency_median"]
    assert body["assumptions"]["eligibility_is_not_hiring"] is True
    assert "ILLUSTRATION ONLY" in body["illustration"]["label"]


def test_impact_estimate_with_no_eligible_users():
    import ml_pipeline.impact as impact
    from tests.test_yojak_api import _market

    body = impact.estimate(_market(), n_users=5, min_pool=10_000)
    assert body["users"] == 0 and body["extra_postings_vs_frequency"]["mean"] is None


# --- acquisition ----------------------------------------------------------------------------------------------

def test_fetch_sends_extra_headers_and_rejects_html_for_pdf(tmp_path, monkeypatch):
    import ml_pipeline.acquire as acquire

    sent = {}

    class Resp(io.BytesIO):
        def __enter__(self):
            return self

        def __exit__(self, *a):
            return False

    def fake_urlopen(req, timeout):
        sent.update(req.headers)
        return Resp(b"%PDF-1.7 ok" if "good" in req.full_url else b"<html>blocked</html>")

    monkeypatch.setattr(acquire.urllib.request, "urlopen", fake_urlopen)
    ok = acquire.fetch("https://x.test/good.pdf", tmp_path / "a.pdf", False, {"Referer": "https://x.test/"})
    assert ok.startswith("downloaded") and sent.get("Referer") == "https://x.test/" and "User-agent" in sent
    bad = acquire.fetch("https://x.test/bad.pdf", tmp_path / "b.pdf", False)
    assert bad.startswith("failed") and not (tmp_path / "b.pdf").exists()
    assert acquire.fetch("https://x.test/good.pdf", tmp_path / "a.pdf", False) == "exists"
    assert all(len(d) in (2, 3) for d in acquire.DOWNLOADS)


# --- evaluation runner ----------------------------------------------------------------------------------------

def test_evaluate_all_runs_selected_steps_in_order_and_reports_failures(monkeypatch):
    import ml_pipeline.evaluate_all as ev

    calls = []

    def fake(step, quick, resume=False):
        calls.append(step)
        if step == "gold":
            raise SystemExit("no labels yet")
        return 0

    monkeypatch.setattr(ev, "run_step", fake)
    assert ev.main(["--only", "docs", "gold", "impact"]) == 1  # gold failed
    assert calls == ["gold", "impact", "docs"]                 # canonical order, not argument order
    calls.clear()
    assert ev.main(["--only", "impact"]) == 0 and calls == ["impact"]


def test_benchmark_checkpoints_resume_and_reset(tmp_path):
    from ml_pipeline.graph.evaluate import Checkpoints

    calls = []

    def work():
        calls.append(1)
        return {"mean": {"t3": 0.5}, "_raw": {"t3": np.arange(3)}}

    first = Checkpoints(tmp_path / "bench", resume=False).run("b1", work)
    again = Checkpoints(tmp_path / "bench", resume=True).run("b1", work)
    assert len(calls) == 1 and again["mean"] == first["mean"] and again["_raw"]["t3"].tolist() == [0, 1, 2]
    # a checkpoint older than the graph data is stale and recomputed
    Checkpoints(tmp_path / "bench", resume=True, newer_than=time.time() + 60).run("b1", work)
    assert len(calls) == 2
    # without --resume the directory is cleared
    Checkpoints(tmp_path / "bench", resume=False).run("b1", work)
    assert len(calls) == 3


# --- hosted-demo snapshot exporter ---------------------------------------------------------------------------

def test_export_static_helpers_make_valid_json():
    import importlib.util
    import json

    spec = importlib.util.spec_from_file_location("export_static", Path(__file__).parents[1] / "scripts" / "export_static.py")
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    assert mod.slug(None) == "all"
    assert mod.slug("Information and Communication Technologies (ICTs)") == "information-and-communication-technologies-icts"
    cleaned = mod.clean({"a": np.float64("nan"), 1: [np.int64(3), float("inf"), (np.bool_(True),)], "b": "x"})
    assert cleaned == {"a": None, "1": [3, None, [True]], "b": "x"}
    json.dumps(cleaned, allow_nan=False)


def test_evidence_status_tracks_which_reports_exist():
    from ml_pipeline.reporting import REPORTS, STATUS_END, STATUS_START, evidence_status, render_readme

    empty = dict.fromkeys(REPORTS)
    assert evidence_status(empty).count("| Pending") >= 5 and "Reading the evidence" not in evidence_status(empty)
    prov = {"provenance": {"generated_at_utc": "2026-10-02T10:00:00+00:00", "git": {"commit": "abcdef123", "dirty": True}}}
    some = {**empty, "data_quality": {**prov, "gold_evaluation": {"skills": {"status": "pending team labels"}}},
            "model_comparison": {**prov, "models": {}, "quick_mode": False}, "upskilling_eval": prov}
    out = evidence_status(some)
    assert "[Graph comparison and ablations](reports/model_comparison.json) | Available" in out
    assert "Partly available" in out                      # upskilling without impact
    assert "`abcdef1`" in out and "uncommitted changes" in out and "2026-10-02" in out
    readme = "\n".join(["a", "<!-- results:start -->", "x", "<!-- results:end -->", "b", STATUS_START, "STALE", STATUS_END, "c"])
    rendered = render_readme(readme, empty)
    assert "STALE" not in rendered and rendered.endswith("c") and render_readme(rendered, empty) == rendered


# --- multilingual extraction: fillers, loanwords, mixed scripts ------------------------------------------------

def test_trim_stopwords_and_hindi_sentence_splitting():
    from app.core.extract import STOPWORDS, candidate_phrases, trim_stopwords

    assert trim_stopwords("मुझे एक्सेल आता है", STOPWORDS) == "एक्सेल"
    assert trim_stopwords("बिक्री का 2 साल का अनुभव", STOPWORDS) == "बिक्री"
    assert trim_stopwords("मुझे है", STOPWORDS) == ""
    text = "मुझे एक्सेल और अकाउंटिंग आती है। थोड़ी अंग्रेज़ी में बात कर सकती हूँ, ग्राहक सेवा का अनुभव है।"
    assert candidate_phrases(text, STOPWORDS) == ["एक्सेल", "अकाउंटिंग", "अंग्रेज़ी", "ग्राहक सेवा"]
    # Gurmukhi "and" ends in a vowel sign, where \b would not see a word boundary
    assert candidate_phrases("ਮੈਨੂੰ ਐਕਸਲ ਅਤੇ ਟੈਲੀ ਆਉਂਦੀ ਹੈ", STOPWORDS) == ["ਐਕਸਲ", "ਟੈਲੀ"]
    assert candidate_phrases("mujhe excel aur tally aata hai", STOPWORDS) == ["excel", "tally"]
    # English behaviour is unchanged when no stopword set is given
    assert candidate_phrases("Python, SQL and standing orders") == ["Python", "SQL", "standing orders"]


def test_loanword_lexicon_is_clean_and_matches_longest_span():
    from app.core.extract import STOPWORDS, load_loanwords, split_loanwords

    raw = pd.read_csv(Path(__file__).parents[1] / "data" / "reference" / "loanwords_hi_pa.csv", comment="#", dtype=str)
    assert not raw["term"].duplicated().any() and set(raw["script"]) == {"hi", "pa"}
    assert raw["english"].str.fullmatch(r"[a-z ]+").all()
    for term, script in zip(raw["term"], raw["script"], strict=True):   # every spelling is in its declared script
        block = range(0x0900, 0x0980) if script == "hi" else range(0x0A00, 0x0A80)
        assert all(ord(ch) in block or ch == " " for ch in term), term
    lex = load_loanwords()
    hits, rest = split_loanwords("मुझे एमएस एक्सेल और डेटा एंट्री का काम", lex, STOPWORDS)
    assert hits == [("एमएस एक्सेल", "ms excel"), ("डेटा एंट्री", "data entry")] and rest == "और"
    assert split_loanwords("ग्राहक सेवा", lex, STOPWORDS) == ([], "ग्राहक सेवा")


def test_extract_routes_loanwords_latin_words_and_native_phrases():
    from app.core.extract import SkillExtractor

    class Linker:
        def __init__(self, table):
            self.table, self.seen = table, []

        def link(self, texts, threshold, single_token_threshold=None, top_k=5):
            self.seen.append(list(texts))
            rows = [{"text": t, "uri": self.table.get(t, (None,))[0], "label": self.table.get(t, (None, None))[1],
                     "score": 0.9 if t in self.table else 0.3, "method": "embedding", "accepted": t in self.table} for t in texts]
            return pd.DataFrame(rows)

    ex = object.__new__(SkillExtractor)
    ex.exact = {"excel": "u:xl", "accounting": "u:acc"}
    ex.label_of = {"u:xl": "use spreadsheets software", "u:acc": "accounting", "u:cs": "customer service"}
    ex.threshold, ex.single_token_threshold, ex.threshold_status, ex.max_n = 0.7, 0.8, "test", 6
    english, multi = Linker({}), Linker({"ग्राहक सेवा": ("u:cs", "customer service")})
    ex._linker = lambda multilingual: multi if multilingual else english

    out = ex.extract("मुझे एक्सेल, टैली और अकाउंटिंग आती है। ग्राहक सेवा का अनुभव है, Excel भी")
    by = {s["source_text"]: s for s in out["skills"]}
    assert out["language"] == "hi"
    assert by["Excel"]["method"] == "exact"                         # Latin word inside Hindi text
    assert by["अकाउंटिंग"]["method"] == "transliteration" and by["अकाउंटिंग"]["uri"] == "u:acc"
    assert by["ग्राहक सेवा"]["uri"] == "u:cs"                       # native phrase, multilingual model
    assert multi.seen == [["ग्राहक सेवा"]]                          # loanwords never reach the multilingual model
    assert english.seen == [["tally"]] and out["unlinked_phrases"] == ["टैली (tally)"]   # honest: no ESCO skill


def test_brute_force_check_runs_on_real_sized_pools():
    from ml_pipeline.upskilling.evaluate import brute_force_check
    from tests.test_yojak_api import _market

    m = _market()
    instances = [{"pool": m.pool(isco2="25"), "have": np.array([2])}, {"pool": m.pool(), "have": np.array([0])}]
    out = brute_force_check(m, instances, n=6, max_cands=4)
    assert out["instances_checked"] == 6 and out["ilp_equals_brute_force"] == 6 and out["ilp_proven_optimal"] == 6
    assert out["candidate_skills_per_instance"] == 4 and out["k_values"] == [1, 2, 3]
